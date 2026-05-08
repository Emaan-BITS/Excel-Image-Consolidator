import os
import zipfile
import shutil
import io
import uuid
from flask import Flask, render_template, request, send_file, jsonify
from werkzeug.utils import secure_filename
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Alignment
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage, ImageDraw

app = Flask(__name__)

# --- GLOBAL PROGRESS TRACKER ---
# Stores the current status of each user's upload session
status_tracker = {}

# Auto-generate a placeholder if it doesn't exist
PLACEHOLDER_IMAGE = "placeholder.jpeg"
if not os.path.exists(PLACEHOLDER_IMAGE):
    img = PILImage.new('RGB', (300, 300), color=(220, 220, 220))
    d = ImageDraw.Draw(img)
    d.text((100, 140), "NO IMAGE", fill=(255, 0, 0))
    img.save(PLACEHOLDER_IMAGE)

# --- NEW STATUS ENDPOINT ---
@app.route("/status/<session_id>")
def get_status(session_id):
    """Returns the current progress for a given session ID."""
    return jsonify(status_tracker.get(session_id, {"progress": 0, "message": "Initializing..."}))

@app.route("/", methods=["GET", "POST"])
def home():
    if request.method == "POST":
        
        # Capture the session ID sent from the frontend, or generate one
        session_id = request.form.get("session_id", str(uuid.uuid4()))
        status_tracker[session_id] = {"progress": 2, "message": "Receiving files and setting up workspace..."}
        
        id_column = request.form.get("column_letter", "A").upper().strip()
        if not id_column.isalpha():
            id_column = "A" 
            
        try:
            target_size = int(request.form.get("image_size", 100))
        except ValueError:
            target_size = 100
            
        dynamic_row_height = target_size * 0.78
        dynamic_img_dim = target_size * 0.9
        dynamic_col_width = (target_size / 100) * 14
            
        excel_files = request.files.getlist("excel_files")
        image_files = request.files.getlist("image_files") 
        
        # Workspace Setup
        workspace_dir = os.path.join("temp_workspace", session_id)
        output_excels_dir = os.path.join(workspace_dir, "outputs")
        temp_img_dir = os.path.join(workspace_dir, "compressed_images")
        raw_img_dir = os.path.join(workspace_dir, "raw_images")
        
        for d in [workspace_dir, output_excels_dir, temp_img_dir, raw_img_dir]:
            os.makedirs(d, exist_ok=True)
            
        # Image Processing
        status_tracker[session_id] = {"progress": 5, "message": "Cataloging uploaded images..."}
        image_map = {}
        for img in image_files:
            if img.filename:
                raw_filename = img.filename.split('/')[-1].split('\\')[-1]
                base_name = os.path.splitext(raw_filename)[0].lower()
                safe_name = secure_filename(raw_filename)
                
                if not safe_name:
                    safe_name = f"img_{uuid.uuid4().hex}.jpg"
                    
                save_path = os.path.join(raw_img_dir, safe_name)
                img.save(save_path)
                image_map[base_name] = save_path

        missing_images_log = []
        total_excels = len([f for f in excel_files if f.filename != '' and f.filename.lower().endswith('.xlsx')])
        current_excel_idx = 0

        # Excel Processing
        for excel_file in excel_files:
            if excel_file.filename == '' or not excel_file.filename.lower().endswith('.xlsx'):
                continue
                
            safe_filename = secure_filename(excel_file.filename)
            input_excel_path = os.path.join(workspace_dir, safe_filename)
            excel_file.save(input_excel_path)
            
            try:
                status_tracker[session_id] = {"progress": 10, "message": f"Opening {safe_filename}..."}
                wb = load_workbook(input_excel_path)
                sheet = wb.active 
                
                empty_col_idx = sheet.max_column + 1
                IMAGE_COLUMN = get_column_letter(empty_col_idx)
                VALUE_COLUMN = get_column_letter(empty_col_idx + 1)
                
                sheet[f"{IMAGE_COLUMN}1"] = "Product Image"
                sheet[f"{VALUE_COLUMN}1"] = "Image Value"

                total_rows = sheet.max_row
                
                for row in range(2, total_rows + 1):
                    
                    # Update Progress every 5 rows to avoid spamming the tracker
                    if row % 5 == 0 or row == total_rows:
                        # Base progress (10%) + distribute the remaining 80% among the files/rows
                        file_base_prog = 10 + (current_excel_idx / total_excels) * 80
                        row_prog = (row / total_rows) * (80 / total_excels)
                        current_prog = int(file_base_prog + row_prog)
                        
                        status_tracker[session_id] = {
                            "progress": current_prog, 
                            "message": f"Compressing & Inserting: Row {row} of {total_rows} ({safe_filename})"
                        }

                    sheet.row_dimensions[row].height = dynamic_row_height 
                    sheet.column_dimensions[IMAGE_COLUMN].width = dynamic_col_width 
                    sheet.column_dimensions[VALUE_COLUMN].width = 15

                    item_id_cell = sheet[f"{id_column}{row}"]
                    item_id_cell.number_format = '@' 
                    
                    if item_id_cell.value is not None:
                        item_id = str(item_id_cell.value).strip()
                        item_id_cell.value = item_id 
                        
                        item_id_lower = item_id.lower()
                        value_cell = sheet[f"{VALUE_COLUMN}{row}"]
                        
                        if item_id_lower in image_map:
                            original_img_path = image_map[item_id_lower]
                            value_cell.value = True
                        else:
                            missing_images_log.append(f"File: {safe_filename} | Row: {row} | ID: {item_id}")
                            original_img_path = PLACEHOLDER_IMAGE
                            value_cell.value = False
                        
                        try:
                            if os.path.exists(original_img_path):
                                optimized_img_path = os.path.join(temp_img_dir, f"opt_{row}_{safe_filename}.jpg")
                                
                                with PILImage.open(original_img_path) as pil_img:
                                    if pil_img.mode in ("RGBA", "P"):
                                        pil_img = pil_img.convert("RGB")
                                    pil_img.thumbnail((300, 300), PILImage.Resampling.LANCZOS)
                                    pil_img.save(optimized_img_path, format="JPEG", optimize=True, quality=75)

                                img = XLImage(optimized_img_path)
                                img.width = dynamic_img_dim
                                img.height = dynamic_img_dim
                                
                                col_idx = column_index_from_string(IMAGE_COLUMN) - 1
                                row_idx = row - 1
                                offset = pixels_to_EMU(2)
                                size = pixels_to_EMU(target_size)
                                
                                marker_from = AnchorMarker(col=col_idx, colOff=offset, row=row_idx, rowOff=offset)
                                marker_to = AnchorMarker(col=col_idx, colOff=offset + size, row=row_idx, rowOff=offset + size)
                                
                                anchor = TwoCellAnchor(editAs='twoCell')
                                anchor._from = marker_from
                                anchor.to = marker_to
                                
                                img.anchor = anchor
                                sheet.add_image(img)
                        except Exception as e:
                            print(f"Error inserting image for row {row} in {safe_filename}: {e}")

                status_tracker[session_id] = {"progress": 90, "message": f"Formatting tables in {safe_filename}..."}
                
                for col_idx in range(1, sheet.max_column + 1):
                    header_cell = sheet.cell(row=1, column=col_idx)
                    if header_cell.value is None or str(header_cell.value).strip() == "":
                        header_cell.value = f"Header_{get_column_letter(col_idx)}"
                    else:
                        header_cell.value = str(header_cell.value)

                center_aligned = Alignment(horizontal='center', vertical='center')
                for row_cells in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
                    for cell in row_cells:
                        cell.alignment = center_aligned

                for existing_table in list(sheet.tables.keys()):
                    del sheet.tables[existing_table]
                    
                final_col_letter = get_column_letter(sheet.max_column)
                table_ref = f"A1:{final_col_letter}{sheet.max_row}"
                tab = Table(displayName=f"Table_{safe_filename.replace('.xlsx', '').replace(' ', '_')}", ref=table_ref)
                style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
                tab.tableStyleInfo = style
                sheet.add_table(tab)

                output_excel_path = os.path.join(output_excels_dir, f"Processed_{safe_filename}")
                wb.save(output_excel_path)
                current_excel_idx += 1
            
            except Exception as e:
                print(f"Failed to process {safe_filename}: {e}")

        # Missing Images Log
        if missing_images_log:
            report_path = os.path.join(output_excels_dir, "Missing_Images_Report.txt")
            with open(report_path, "w") as f:
                f.write("MISSING IMAGES REPORT\n")
                f.write("="*40 + "\n")
                f.write("The following Product IDs did not have a matching uploaded image.\n\n")
                for entry in missing_images_log:
                    f.write(entry + "\n")

        # Zipping Files
        status_tracker[session_id] = {"progress": 95, "message": "Zipping files for download..."}
        memory_file = io.BytesIO()
        if os.listdir(output_excels_dir):
            with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as out_zip:
                for file in os.listdir(output_excels_dir):
                    out_zip.write(os.path.join(output_excels_dir, file), file)
            memory_file.seek(0)
            
            try:
                shutil.rmtree(workspace_dir)
            except Exception as e:
                print(f"Housekeeping failed: {e}")

            status_tracker[session_id] = {"progress": 100, "message": "Done! Your download is starting."}
            return send_file(memory_file, download_name="Finished_Excel_Files.zip", as_attachment=True)
        else:
            shutil.rmtree(workspace_dir)
            status_tracker[session_id] = {"progress": 0, "message": "Failed: No valid files."}
            return "No valid .xlsx files were uploaded or processed.", 400

    return render_template("index.html")

if __name__ == "__main__":
    app.run(host='0.0.0.0', port=5000, debug=True)